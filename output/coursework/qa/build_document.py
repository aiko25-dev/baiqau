from pathlib import Path
import re, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'Baiqau_tusindirme_zhazba.txt'
target = ROOT / 'Baiqau_tusindirme_zhazba.docx'
lines = source.read_text(encoding='utf-8').splitlines()
doc = Document()

def font(style, size=14):
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.bold=False
    style.font.italic=False
    style.font.underline=False
    style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts'); rpr.insert(0,rf)
    for attr in list(rf.attrib):
        if 'Theme' in attr:
            del rf.attrib[attr]
    for k in ('ascii','hAnsi','eastAsia','cs'):
        rf.set(qn('w:'+k),'Times New Roman')
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'kk-KZ'); rpr.append(lang)

def layout(sec):
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.left_margin=Cm(3); sec.right_margin=Cm(1)
    sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
    sec.footer_distance=Cm(1)

for sec in doc.sections: layout(sec)
for name in ('Normal','Heading 1','Heading 2','Title','Caption'):
    s=doc.styles[name]; font(s)
    f=s.paragraph_format
    f.first_line_indent=Cm(1.25)
    f.line_spacing=1.5
    f.space_before=Pt(0); f.space_after=Pt(0)
    f.widow_control=True
    f.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

for name in ('Heading 1','Heading 2'):
    f=doc.styles[name].paragraph_format
    f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.keep_with_next=True
    f.space_after=Pt(21)
    if name=='Heading 2': f.space_before=Pt(21)

doc.core_properties.title='Байқаулар өткізуге және қатысушылардың жұмыстарын бағалауға арналған мобильді платформа құру'
doc.core_properties.subject='Baiqau мобильді қосымшасының түсіндірме жазбасы'
doc.core_properties.author=''
doc.core_properties.keywords='Baiqau, Flutter, Dart, байқау, бағалау'
settings=doc.settings.element
upd=OxmlElement('w:updateFields'); upd.set(qn('w:val'),'true'); settings.append(upd)

majors={'КІРІСПЕ','1 БАЙҚАУЛАРДЫ ҰЙЫМДАСТЫРУ ЖӘНЕ БАҒАЛАУ ЖҮЙЕЛЕРІНІҢ ТЕОРИЯЛЫҚ НЕГІЗДЕРІ','2 МОБИЛЬДІ ПЛАТФОРМАНЫ ЖОБАЛАУ','3 МОБИЛЬДІ ҚОСЫМШАНЫ ІСКЕ АСЫРУ ЖӘНЕ ТЕСТІЛЕУ','ҚОРЫТЫНДЫ','ПАЙДАЛАНЫЛҒАН ӘДЕБИЕТТЕР ТІЗІМІ','ҚОСЫМШАЛАР'}
toc=True
table_no=0
i=0
while i<len(lines):
    line=lines[i]
    if not line:
        i+=1; continue
    if line=='МАЗМҰНЫ':
        p=doc.add_paragraph(line,'Heading 1'); i+=1; continue
    if line=='КІРІСПЕ' and toc:
        toc=False
        sec=doc.add_section(WD_SECTION_START.NEW_PAGE); layout(sec)
        sec.footer.is_linked_to_previous=False
        p=sec.footer.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=1
        r=p.add_run(); r.font.name='Times New Roman'; r.font.size=Pt(14)
        fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
    elif line in majors:
        p=doc.add_paragraph(line,'Heading 1')
        p.paragraph_format.page_break_before=True
        i+=1; continue
    if line=='КІРІСПЕ':
        doc.add_paragraph(line,'Heading 1'); i+=1; continue
    if toc:
        p=doc.add_paragraph(line.replace('\t','\t '))
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        i+=1; continue
    if re.match(r'^[123]\.[1-4] ',line):
        doc.add_paragraph(line,'Heading 2'); i+=1; continue
    if line in {'Қосымша А','Қосымша Ә'}:
        p=doc.add_paragraph(line,'Heading 2')
        if line=='Қосымша Ә': p.paragraph_format.page_break_before=True
        i+=1; continue
    if line.startswith('Кесте '):
        p=doc.add_paragraph(line)
        p.paragraph_format.keep_with_next=True
        table_no+=1
        rows=[]
        i+=1
        while i<len(lines) and '\t' in lines[i]:
            rows.append(lines[i].split('\t')); i+=1
        if not rows: raise ValueError('Missing table '+line)
        widths={3:[3.8,6.2,7.0],4:[1.6,6.0,4.4,5.0]}.get(len(rows[0]),[17/len(rows[0])]*len(rows[0]))
        if line.startswith('Кесте 1.2'): widths=[5.7,3.0,3.0,5.3]
        if line.startswith('Кесте 2.1'): widths=[1.6,6.2,9.2]
        if line.startswith('Кесте 2.3'): widths=[3.4,7.4,6.2]
        if line.startswith('Кесте 3.2'): widths=[3.4,5.0,4.4,4.2]
        chunks=[rows]
        if line.startswith('Кесте 3.1'):
            chunks=[rows[:11],[rows[0],*rows[11:]]]
        for chunk_index,chunk in enumerate(chunks):
            if chunk_index:
                caption=doc.add_paragraph('3.1-кестенің жалғасы')
                caption.paragraph_format.keep_with_next=True
            table=doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit=False
            table.style='Table Grid'
            for col,w in zip(table.columns,widths): col.width=Cm(w)
            for ri,row in enumerate(chunk):
                cells=table.add_row().cells
                for ci,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[ci])
                    p=cell.paragraphs[0]
                    p.paragraph_format.first_line_indent=Pt(0)
                    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.line_spacing=1.5
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.space_before=Pt(0)
                    p.paragraph_format.keep_with_next=ri<len(chunk)-1
                    r=p.add_run(value); r.font.name='Times New Roman'; r.font.size=Pt(12)
                    r.font.bold=False
                trpr=table.rows[-1]._tr.get_or_add_trPr()
                no=OxmlElement('w:cantSplit'); trpr.append(no)
                if ri==0:
                    repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
        continue
    if line.startswith('[ОСЫ ЖЕРГЕ СУРЕТ ҚОЮ:'):
        p=doc.add_paragraph(line)
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(7)
        i+=1; continue
    if line.startswith('Сурет '):
        p=doc.add_paragraph(line,'Caption')
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(7)
        i+=1; continue
    if line.startswith('S = ') or line.startswith('R = '):
        # Native editable Office Math with subscripted terms and a fraction.
        p=doc.add_paragraph()
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0)
        math=OxmlElement('m:oMath')
        def mr(parent,s):
            r=OxmlElement('m:r'); t=OxmlElement('m:t'); t.text=s; r.append(t); parent.append(r)
        def sub(parent,s,n):
            ss=OxmlElement('m:sSub')
            e=OxmlElement('m:e'); mr(e,s)
            sb=OxmlElement('m:sub'); mr(sb,n)
            ss.extend([e,sb]); parent.append(ss)
        if line.startswith('S = '):
            mr(math,'S = (0,40 × C + 0,35 × Q + 0,25 × P) × 10')
            number='(2.1)'
        else:
            mr(math,'R = ')
            f=OxmlElement('m:f'); num=OxmlElement('m:num'); den=OxmlElement('m:den')
            sub(num,'S','1'); mr(num,' + '); sub(num,'S','2'); mr(num,' + … + '); sub(num,'S','n'); mr(den,'n')
            f.extend([num,den]); math.append(f); mr(math,',  n > 0')
            number='(2.2)'
        p._p.append(math); p.add_run('   '+number)
        i+=1; continue
    p=doc.add_paragraph(line)
    if line.startswith(('static double','  ','}','.\\scripts\\','No issues found!','00:06')):
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs: r.font.size=Pt(12)
    if re.match(r'^\d+\. ',line):
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_together=True
    i+=1

doc.save(target)
print(json.dumps({'docx':str(target),'words':len(source.read_text(encoding='utf-8').split()),'tables':table_no},ensure_ascii=False))
